"""Convert docs/training-manual.md to lesotho-vitallink-docs-v2.docx using python-docx."""
import re
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REPO = Path(__file__).parent.parent
MD   = REPO / "docs" / "training-manual.md"
OUT  = REPO / "lesotho-vitallink-docs-v2.docx"

# ── helpers ──────────────────────────────────────────────────────────────────

def set_heading_color(para, hex_color):
    for run in para.runs:
        run.font.color.rgb = RGBColor.from_string(hex_color)

def add_horizontal_rule(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'AAAAAA')
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_after = Pt(6)

def apply_inline(run_text, para):
    """Split text on inline code/bold/italic markers and add styled runs."""
    # patterns: **bold**, *italic*, `code`
    parts = re.split(r'(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)', run_text)
    for part in parts:
        if part.startswith('`') and part.endswith('`'):
            r = para.add_run(part[1:-1])
            r.font.name = 'Courier New'
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        elif part.startswith('**') and part.endswith('**'):
            r = para.add_run(part[2:-2])
            r.bold = True
        elif part.startswith('*') and part.endswith('*'):
            r = para.add_run(part[1:-1])
            r.italic = True
        else:
            para.add_run(part)

def add_para(doc, text, style='Normal', indent=None):
    p = doc.add_paragraph(style=style)
    apply_inline(text, p)
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    return p

def set_col_widths(table, widths_cm):
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths_cm):
                cell.width = Cm(widths_cm[i])

# ── styles ───────────────────────────────────────────────────────────────────

def setup_styles(doc):
    styles = doc.styles

    normal = styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)

    for level, size, bold in [
        ('Heading 1', 18, True),
        ('Heading 2', 14, True),
        ('Heading 3', 12, True),
        ('Heading 4', 11, True),
    ]:
        s = styles[level]
        s.font.name = 'Calibri'
        s.font.size = Pt(size)
        s.font.bold = bold
        s.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
        s.paragraph_format.space_before = Pt(12)
        s.paragraph_format.space_after  = Pt(4)

    # Code block style (based on Normal)
    if 'Code Block' not in [s.name for s in styles]:
        code_style = styles.add_style('Code Block', 1)
        code_style.base_style = styles['Normal']
    cs = styles['Code Block']
    cs.font.name = 'Courier New'
    cs.font.size = Pt(8.5)
    cs.paragraph_format.left_indent  = Inches(0.3)
    cs.paragraph_format.space_before = Pt(2)
    cs.paragraph_format.space_after  = Pt(2)

# ── main parser ──────────────────────────────────────────────────────────────

def convert(md_path, out_path):
    doc = Document()
    setup_styles(doc)

    # Page margins
    for section in doc.sections:
        section.top_margin    = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    lines = md_path.read_text(encoding='utf-8').splitlines()

    in_code   = False
    code_lang = ''
    in_table  = False
    table_obj = None
    table_rows = []
    i = 0

    while i < len(lines):
        line = lines[i]

        # ── fenced code block ──
        if line.startswith('```'):
            if not in_code:
                in_code   = True
                code_lang = line[3:].strip()
            else:
                in_code = False
                code_lang = ''
            i += 1
            continue

        if in_code:
            p = doc.add_paragraph(line, style='Code Block')
            # light grey background via paragraph shading
            pPr = p._p.get_or_add_pPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), 'F3F3F3')
            pPr.append(shd)
            i += 1
            continue

        # ── horizontal rule ──
        if re.match(r'^-{3,}$', line.strip()):
            add_horizontal_rule(doc)
            i += 1
            continue

        # ── headings ──
        m = re.match(r'^(#{1,4})\s+(.*)', line)
        if m:
            level = len(m.group(1))
            text  = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', m.group(2))  # strip links
            style = f'Heading {level}'
            doc.add_heading(text, level=level)
            i += 1
            continue

        # ── table ──
        if line.startswith('|'):
            if not in_table:
                in_table   = True
                table_rows = []
            # skip separator rows (|---|---|)
            if re.match(r'^\|[\s\-\|:]+\|$', line):
                i += 1
                continue
            cells = [c.strip() for c in line.strip('|').split('|')]
            table_rows.append(cells)
            i += 1
            # peek: if next line is not a table row, flush
            if i >= len(lines) or not lines[i].startswith('|'):
                in_table = False
                if table_rows:
                    ncols = max(len(r) for r in table_rows)
                    tbl = doc.add_table(rows=len(table_rows), cols=ncols)
                    tbl.style = 'Table Grid'
                    for ri, row_data in enumerate(table_rows):
                        for ci, cell_text in enumerate(row_data):
                            cell = tbl.cell(ri, ci)
                            cell.text = ''
                            p = cell.paragraphs[0]
                            apply_inline(cell_text, p)
                            p.paragraph_format.space_after = Pt(2)
                            if ri == 0:
                                for run in p.runs:
                                    run.bold = True
                                # header shading
                                tc_pr = cell._tc.get_or_add_tcPr()
                                shd = OxmlElement('w:shd')
                                shd.set(qn('w:val'), 'clear')
                                shd.set(qn('w:color'), 'auto')
                                shd.set(qn('w:fill'), 'D6E4F0')
                                tc_pr.append(shd)
                    doc.add_paragraph()  # spacing after table
                table_rows = []
            continue

        in_table = False

        # ── blockquote ──
        if line.startswith('>'):
            text = line.lstrip('> ').strip()
            p = add_para(doc, text, indent=0.4)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after  = Pt(2)
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue

        # ── list items ──
        m = re.match(r'^(\s*)([-*+]|\d+\.)\s+(.*)', line)
        if m:
            indent_spaces = len(m.group(1))
            text  = m.group(3)
            level = indent_spaces // 2
            style = 'List Bullet' if re.match(r'[-*+]', m.group(2)) else 'List Number'
            p = add_para(doc, text, style=style, indent=level * 0.2 if level else None)
            i += 1
            continue

        # ── blank line ──
        if not line.strip():
            doc.add_paragraph()
            i += 1
            continue

        # ── normal paragraph ──
        add_para(doc, line)
        i += 1

    doc.save(out_path)
    print(f"Saved: {out_path}")

if __name__ == '__main__':
    md_arg  = Path(sys.argv[1]) if len(sys.argv) > 1 else MD
    out_arg = Path(sys.argv[2]) if len(sys.argv) > 2 else OUT
    convert(md_arg, out_arg)
